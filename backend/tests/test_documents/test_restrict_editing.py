"""Tests for app.documents.restrict_editing.

The salt (and therefore the hash) is random on every call, so these tests
assert structure and lengths, never specific hash values -- see the
signatures feature plan, Task 2 determinism note.
"""
import base64
import io

from docx import Document
from docx.oxml.ns import qn

from app.documents.restrict_editing import apply_read_only_protection, save_docx


def _protection_node(document: Document):
    settings = document.settings.element
    nodes = settings.findall(qn("w:documentProtection"))
    assert len(nodes) == 1, f"expected exactly one w:documentProtection, found {len(nodes)}"
    return nodes[0]


class TestApplyReadOnlyProtection:
    def test_applies_expected_attributes(self):
        document = Document()
        apply_read_only_protection(document, "correct-horse-battery-staple")

        node = _protection_node(document)
        assert node.get(qn("w:edit")) == "readOnly"
        assert node.get(qn("w:enforcement")) == "1"
        assert node.get(qn("w:cryptProviderType")) == "rsaAES"
        assert node.get(qn("w:cryptAlgorithmClass")) == "hash"
        assert node.get(qn("w:cryptAlgorithmType")) == "typeAny"
        assert node.get(qn("w:cryptAlgorithmSid")) == "14"
        assert node.get(qn("w:cryptSpinCount")) == "100000"

        hash_bytes = base64.b64decode(node.get(qn("w:hash")))
        salt_bytes = base64.b64decode(node.get(qn("w:salt")))
        assert len(hash_bytes) == 64  # SHA-512 digest
        assert len(salt_bytes) == 16

    def test_is_first_child_of_settings(self):
        document = Document()
        apply_read_only_protection(document, "password")

        settings = document.settings.element
        assert settings[0].tag == qn("w:documentProtection")

    def test_survives_save_and_reopen(self):
        document = Document()
        apply_read_only_protection(document, "password")
        output_bytes = save_docx(document)

        reopened = Document(io.BytesIO(output_bytes))
        node = _protection_node(reopened)
        assert node.get(qn("w:edit")) == "readOnly"

    def test_reapplication_is_idempotent(self):
        document = Document()
        apply_read_only_protection(document, "first-password")
        apply_read_only_protection(document, "second-password")

        # Exactly one node, not two -- and it reflects the second call.
        settings = document.settings.element
        assert len(settings.findall(qn("w:documentProtection"))) == 1

    def test_different_calls_produce_different_salts(self):
        document_a = Document()
        document_b = Document()
        apply_read_only_protection(document_a, "same-password")
        apply_read_only_protection(document_b, "same-password")

        salt_a = _protection_node(document_a).get(qn("w:salt"))
        salt_b = _protection_node(document_b).get(qn("w:salt"))
        assert salt_a != salt_b


class TestSaveDocx:
    def test_returns_bytes_reopenable_as_docx(self):
        document = Document()
        document.add_paragraph("hello")
        output_bytes = save_docx(document)

        assert isinstance(output_bytes, bytes)
        reopened = Document(io.BytesIO(output_bytes))
        assert reopened.paragraphs[0].text == "hello"
