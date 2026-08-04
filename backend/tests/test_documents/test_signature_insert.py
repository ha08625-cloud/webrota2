"""Tests for app.documents.signature_insert.insert_signature.

No DB, no fixtures on disk: every docx used here is built in-test with
python-docx and round-tripped through bytes, matching the module under
test's own bytes-in/bytes-out contract.
"""
import base64
import io

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.documents.errors import DocumentFormatError
from app.documents.signature_insert import insert_date, insert_signature

# A minimal valid 1x1 red PNG -- small enough to hardcode, but a real,
# structurally valid PNG (verified: decodes to a 1x1 image).
MINIMAL_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def _minimal_png_bytes() -> bytes:
    return base64.b64decode(MINIMAL_PNG_B64)


def _make_docx_bytes(rows: int = 3, cols: int = 2) -> bytes:
    """Builds a docx with one table (rows x cols), some text in the
    non-target cells, and returns it as bytes -- mirroring the real
    signature-block template shape (2 columns, signature cell is the
    last row's first cell, left deliberately empty)."""
    document = Document()
    table = document.add_table(rows=rows, cols=cols)
    for r in range(rows):
        for c in range(cols):
            if not (r == rows - 1 and c == 0):
                table.rows[r].cells[c].text = f"r{r}c{c}"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _image_part_count(document: Document) -> int:
    return sum(
        1 for part in document.part.related_parts.values()
        if getattr(part, "content_type", "").startswith("image/")
    )


class TestInsertSignature:
    def test_inserts_image_into_bottom_left_cell(self):
        docx_bytes = _make_docx_bytes()
        image_bytes = _minimal_png_bytes()

        before_count = _image_part_count(Document(io.BytesIO(docx_bytes)))

        document = insert_signature(docx_bytes, image_bytes)

        buffer = io.BytesIO()
        document.save(buffer)
        reopened = Document(io.BytesIO(buffer.getvalue()))

        target_cell = reopened.tables[0].rows[-1].cells[0]
        drawings = target_cell._element.findall(".//" + qn("w:drawing"))
        assert len(drawings) == 1

        after_count = _image_part_count(reopened)
        assert after_count == before_count + 1

    def test_other_cells_are_untouched(self):
        docx_bytes = _make_docx_bytes(rows=3, cols=2)
        document = insert_signature(docx_bytes, _minimal_png_bytes())

        table = document.tables[0]
        for r in range(3):
            for c in range(2):
                if r == 2 and c == 0:
                    continue
                assert table.rows[r].cells[c].text == f"r{r}c{c}"

    def test_random_bytes_raise_document_format_error(self):
        with pytest.raises(DocumentFormatError):
            insert_signature(b"not a docx file at all", _minimal_png_bytes())

    def test_no_table_raises_document_format_error(self):
        document = Document()
        document.add_paragraph("no tables here")
        buffer = io.BytesIO()
        document.save(buffer)

        with pytest.raises(DocumentFormatError, match="no table"):
            insert_signature(buffer.getvalue(), _minimal_png_bytes())

    def test_wrong_column_count_raises_document_format_error_naming_layout(self):
        docx_bytes = _make_docx_bytes(rows=3, cols=3)

        with pytest.raises(DocumentFormatError, match="2 columns"):
            insert_signature(docx_bytes, _minimal_png_bytes())


class TestInsertDate:
    def test_appends_a_new_paragraph_to_the_bottom_right_cell(self):
        docx_bytes = _make_docx_bytes()
        document = Document(io.BytesIO(docx_bytes))

        insert_date(document, "04/08/2026")

        cell = document.tables[0].rows[-1].cells[1]
        # The cell's existing content ("r2c1", per _make_docx_bytes) is left
        # in place -- the date lands in a new paragraph after it, not
        # overwriting it.
        assert cell.paragraphs[0].text == "r2c1"
        assert cell.paragraphs[-1].text == "04/08/2026"

    def test_other_cells_are_untouched(self):
        docx_bytes = _make_docx_bytes(rows=3, cols=2)
        document = Document(io.BytesIO(docx_bytes))

        insert_date(document, "04/08/2026")

        table = document.tables[0]
        for r in range(3):
            for c in range(2):
                if r == 2 and c == 1:
                    continue
                expected = "" if (r == 2 and c == 0) else f"r{r}c{c}"
                assert table.rows[r].cells[c].text == expected

    def test_survives_a_save_and_reopen_round_trip(self):
        docx_bytes = _make_docx_bytes()
        document = Document(io.BytesIO(docx_bytes))
        insert_date(document, "04/08/2026")

        buffer = io.BytesIO()
        document.save(buffer)
        reopened = Document(io.BytesIO(buffer.getvalue()))

        cell = reopened.tables[0].rows[-1].cells[1]
        assert "04/08/2026" in cell.text
