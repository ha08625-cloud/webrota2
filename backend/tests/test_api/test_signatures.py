"""Signatures router tests (signatures feature, Task 3).

Upload responses use a uniform 200 for both create and replace -- see
schemas/signature.py.

Auth coverage for this router (and every other router) now lives centrally
in test_auth.py (auth plan, Task 4) -- the per-router TestAuth class that
used to live here tested the M3.5 API_TOKEN shim, which get_current_user no
longer implements, and was removed rather than rewritten.

The exception is TestReadAccess at the foot of this file. The two GETs
here are the only reads in the API that are not open to every tier, so
the central sweeps -- which assert the general rule -- cannot carry them.
"""
import base64
import datetime
import io
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.documents.errors import ConversionError
from app.models import Doctor
from app.models.enums import AccessLevel, DoctorType
from app.models.permissions import PRESET_FOR_ACCESS_LEVEL, preset
from tests.soffice_support import requires_soffice

SAMPLE_RTF_PATH = Path(__file__).parent.parent / "fixtures" / "certificate_sample.rtf"

# Minimal valid 1x1 transparent PNG.
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY"
    "42YAAAAASUVORK5CYII="
)


def _make_doctor(db_session, code="CC"):
    doctor = Doctor(code=code, doctor_type=DoctorType.PARTNER, sessions_per_week=10, active=True)
    db_session.add(doctor)
    db_session.commit()
    db_session.refresh(doctor)
    return doctor.id


def _build_docx(with_table=True, cols=2) -> bytes:
    document = Document()
    if with_table:
        document.add_table(rows=2, cols=cols)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestUpload:
    def test_upload_creates_and_returns_meta(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        resp = client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["doctor_id"] == doctor_id
        assert body["content_type"] == "image/png"
        assert "uploaded_at" in body

    def test_upload_replace_updates_timestamp(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        first = client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        ).json()
        second = client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig2.png", _TINY_PNG, "image/png")},
        ).json()
        assert second["uploaded_at"] >= first["uploaded_at"]
        assert len(client.get("/api/v1/signatures").json()) == 1

    def test_wrong_content_type_422(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        resp = client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.gif", b"gifbytes", "image/gif")},
        )
        assert resp.status_code == 422

    def test_unknown_doctor_404(self, client, db_session):
        resp = client.post(
            "/api/v1/signatures/999999",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        )
        assert resp.status_code == 404

    def test_oversize_image_413(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        big = b"\x00" * (5 * 1024 * 1024 + 1)
        resp = client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", big, "image/png")},
        )
        assert resp.status_code == 413


class TestListAndFetch:
    def test_list_and_image_roundtrip(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        )
        listing = client.get("/api/v1/signatures").json()
        assert len(listing) == 1
        assert listing[0]["doctor_id"] == doctor_id

        img_resp = client.get(f"/api/v1/signatures/{doctor_id}/image")
        assert img_resp.status_code == 200
        assert img_resp.headers["content-type"] == "image/png"
        assert img_resp.content == _TINY_PNG

    def test_image_404_when_absent(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        resp = client.get(f"/api/v1/signatures/{doctor_id}/image")
        assert resp.status_code == 404


class TestDelete:
    def test_delete_then_image_404(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        )
        del_resp = client.delete(f"/api/v1/signatures/{doctor_id}")
        assert del_resp.status_code == 204
        assert client.get(f"/api/v1/signatures/{doctor_id}/image").status_code == 404

    def test_delete_when_absent_404(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        resp = client.delete(f"/api/v1/signatures/{doctor_id}")
        assert resp.status_code == 404


class TestApply:
    def test_apply_happy_path(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        )
        docx_bytes = _build_docx()
        resp = client.post(
            f"/api/v1/signatures/{doctor_id}/apply",
            files={
                "file": (
                    "letter.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert "-signed.docx" in resp.headers["content-disposition"]

        out_doc = Document(io.BytesIO(resp.content))
        bottom_left_cell = out_doc.tables[0].rows[-1].cells[0]
        drawings = bottom_left_cell._tc.findall(".//" + qn("w:drawing"))
        assert len(drawings) == 1

        today = datetime.date.today().strftime("%d/%m/%Y")
        bottom_right_cell = out_doc.tables[0].rows[-1].cells[1]
        assert today in bottom_right_cell.text

        protections = out_doc.settings.element.findall(qn("w:documentProtection"))
        assert len(protections) == 1

    def test_apply_no_signature_stored_409(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        docx_bytes = _build_docx()
        resp = client.post(
            f"/api/v1/signatures/{doctor_id}/apply",
            files={
                "file": (
                    "letter.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert resp.status_code == 409
        assert resp.json()["detail"] == "No signature stored for this doctor"

    def test_apply_garbage_bytes_422(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        )
        resp = client.post(
            f"/api/v1/signatures/{doctor_id}/apply",
            files={"file": ("letter.docx", b"not a real docx", "application/octet-stream")},
        )
        assert resp.status_code == 422
        assert "not a valid" in resp.json()["detail"].lower()

    def test_apply_no_table_422(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        )
        docx_bytes = _build_docx(with_table=False)
        resp = client.post(
            f"/api/v1/signatures/{doctor_id}/apply",
            files={
                "file": (
                    "letter.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert resp.status_code == 422
        assert "no table" in resp.json()["detail"].lower()

    def test_apply_oversize_document_413(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        )
        big = b"\x00" * (10 * 1024 * 1024 + 1)
        resp = client.post(
            f"/api/v1/signatures/{doctor_id}/apply",
            files={"file": ("letter.docx", big, "application/octet-stream")},
        )
        assert resp.status_code == 413

    def test_apply_pdf_upload_422(self, client, db_session):
        """A PDF is neither a zip nor RTF, so it never reaches either
        transform -- the sniff rejects it."""
        doctor_id = _make_doctor(db_session)
        client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        )
        resp = client.post(
            f"/api/v1/signatures/{doctor_id}/apply",
            files={"file": ("cert.pdf", b"%PDF-1.7\n%stub\n", "application/pdf")},
        )
        assert resp.status_code == 422
        assert "not a valid" in resp.json()["detail"].lower()


class TestApplyRtf:
    """The .rtf path: splice a signature in, return a PDF."""

    def _prepare(self, client, db_session):
        doctor_id = _make_doctor(db_session)
        client.post(
            f"/api/v1/signatures/{doctor_id}",
            files={"file": ("sig.png", _TINY_PNG, "image/png")},
        )
        return doctor_id

    def _post_rtf(self, client, doctor_id, rtf_bytes, filename="certificate.rtf"):
        return client.post(
            f"/api/v1/signatures/{doctor_id}/apply",
            files={"file": (filename, rtf_bytes, "application/rtf")},
        )

    @requires_soffice
    def test_apply_rtf_returns_pdf(self, client, db_session):
        doctor_id = self._prepare(client, db_session)

        resp = self._post_rtf(client, doctor_id, SAMPLE_RTF_PATH.read_bytes())

        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"] == "application/pdf"
        assert "-signed.pdf" in resp.headers["content-disposition"]
        assert resp.content.startswith(b"%PDF-")

    @requires_soffice
    def test_apply_rtf_ignores_the_declared_content_type(self, client, db_session):
        """Format comes from the bytes, not the browser's guess or the
        extension: RTF content announced as docx still yields a PDF."""
        doctor_id = self._prepare(client, db_session)

        resp = client.post(
            f"/api/v1/signatures/{doctor_id}/apply",
            files={
                "file": (
                    "certificate.docx",
                    SAMPLE_RTF_PATH.read_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"] == "application/pdf"

    def test_apply_rtf_without_the_anchor_422(self, client, db_session):
        """Runs everywhere: the splice fails before any conversion, so this
        needs no LibreOffice."""
        doctor_id = self._prepare(client, db_session)
        without_anchor = SAMPLE_RTF_PATH.read_bytes().replace(
            b"Signature", b"Sign-ature"
        )

        resp = self._post_rtf(client, doctor_id, without_anchor)

        assert resp.status_code == 422
        assert "signature label" in resp.json()["detail"].lower()

    def test_apply_rtf_without_the_date_anchor_422(self, client, db_session):
        """Runs everywhere: same as the Signature-anchor test above, for
        the Date anchor, and needs no LibreOffice either -- the splice
        fails before conversion is attempted."""
        doctor_id = self._prepare(client, db_session)
        without_anchor = SAMPLE_RTF_PATH.read_bytes().replace(b"Date", b"Da-te")

        resp = self._post_rtf(client, doctor_id, without_anchor)

        assert resp.status_code == 422
        assert "date label" in resp.json()["detail"].lower()

    def test_apply_rtf_no_signature_stored_409_before_converting(
        self, client, db_session, monkeypatch
    ):
        """The 409 must come first, so a missing signature never spawns a
        LibreOffice subprocess."""
        doctor_id = _make_doctor(db_session)

        def _fail(*args, **kwargs):
            raise AssertionError("convert_to_pdf must not be called")

        monkeypatch.setattr("app.api.routers.signatures.convert_to_pdf", _fail)

        resp = self._post_rtf(client, doctor_id, SAMPLE_RTF_PATH.read_bytes())

        assert resp.status_code == 409
        assert resp.json()["detail"] == "No signature stored for this doctor"

    def test_conversion_failure_502(self, client, db_session, monkeypatch):
        """A converter failure is ours, not the admin's, and
        the converter's own text stays in the log."""
        doctor_id = self._prepare(client, db_session)

        def _boom(*args, **kwargs):
            raise ConversionError("soffice exited 1: /tmp/xyz/document.rtf")

        monkeypatch.setattr("app.api.routers.signatures.convert_to_pdf", _boom)

        resp = self._post_rtf(client, doctor_id, SAMPLE_RTF_PATH.read_bytes())

        assert resp.status_code == 502
        detail = resp.json()["detail"]
        assert "Could not convert this document to PDF" in detail
        assert "/tmp" not in detail
        assert "soffice" not in detail


class TestReadAccess:
    """Reading a signature needs the `signatures` permission, exactly like
    uploading one.

    The signature image is the one asset in this API worth more outside it
    than in. `signatures` is a boolean rather than a level precisely so
    that there is no read-only view of it to leak through: the people who
    may look at a scanned signature are the people who may upload one, and
    nobody else -- see the router docstring and deps.require_access.

    Note who that now excludes. A Rota admin login writes both rota
    sections and cannot read these two endpoints; under the old tier model
    the same person could. That narrowing is the point of the feature
    rather than a regression, and it is asserted here so that undoing it
    takes a deliberate edit.

    Assertions are `== 403`, never "not 2xx": a broken gate would answer
    404 on the image path for a doctor with nothing stored, and "not 2xx"
    would accept that as a pass.
    """

    def test_a_read_only_login_cannot_list_signatures(self, readonly_client):
        resp = readonly_client.get("/api/v1/signatures")
        assert resp.status_code == 403, resp.text

    def test_a_read_only_login_cannot_fetch_an_image(self, readonly_client):
        resp = readonly_client.get("/api/v1/signatures/1/image")
        assert resp.status_code == 403, resp.text

    def test_the_doctor_tier_preset_cannot_fetch_an_image(
        self, client_with_permissions
    ):
        """The tier most likely to be given a login for its own sake, and
        the one this gate exists to keep out. Asserted through the preset
        its label maps to, since the label itself is not consulted."""
        client = client_with_permissions(
            preset(PRESET_FOR_ACCESS_LEVEL[AccessLevel.DOCTOR.value])
        )
        assert client.get("/api/v1/signatures/1/image").status_code == 403

    def test_a_rota_admin_cannot_list_signatures(self, rota_admin_client):
        """Rota write access is not signature access any more."""
        assert rota_admin_client.get("/api/v1/signatures").status_code == 403

    def test_a_rota_admin_cannot_fetch_a_signature_image(self, rota_admin_client):
        assert rota_admin_client.get("/api/v1/signatures/1/image").status_code == 403

    def test_a_documents_login_can_list_signatures(self, documents_client):
        """No rota access at all, but `signatures` -- the login the
        permission split exists to make possible."""
        resp = documents_client.get("/api/v1/signatures")
        assert resp.status_code == 200, resp.text
        assert resp.json() == []

    def test_gets_the_ordinary_404_for_a_missing_image(self, manager_client):
        """Past the gate, the endpoint behaves as it always did."""
        assert manager_client.get("/api/v1/signatures/1/image").status_code == 404

    def test_manager_can_list_signatures(self, manager_client):
        assert manager_client.get("/api/v1/signatures").status_code == 200
