"""POST /eoi/fill (study EOI autofill, Task 2).

The fill itself is covered in tests/test_documents/test_eoi_fill.py; what
matters here is the HTTP envelope -- the upload guards, the reported
misses, and the filename echoed back.
"""
import io
from pathlib import Path

import pytest
from docx import Document

BLANK_FORM = Path(__file__).parents[1] / "fixtures" / "site_id_form_blank.docx"

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


@pytest.fixture(scope="module")
def blank() -> bytes:
    return BLANK_FORM.read_bytes()


def _unrelated_docx() -> bytes:
    document = Document()
    document.add_paragraph("Nothing to do with site identification.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _post(client, filename, data, content_type="application/octet-stream"):
    return client.post(
        "/api/v1/eoi/fill", files={"file": (filename, data, content_type)}
    )


class TestFill:
    def test_happy_path_returns_a_docx(self, client, blank):
        resp = _post(client, "site id form.docx", blank)
        assert resp.status_code == 200
        assert resp.headers["content-type"] == DOCX_MEDIA_TYPE
        # A real, openable package, not the bytes we sent.
        filled = Document(io.BytesIO(resp.content))
        assert resp.content != blank
        text = "\n".join(
            cell.text
            for table in filled.tables
            for row in table.rows
            for cell in row.cells
        )
        assert "Summertown Health Centre" in text

    def test_only_section_10_is_unmatched_on_the_blank_form(self, client, blank):
        # The blank form has no "Non-Commercial ... Studies" cell, so that
        # one rule reports a miss and the other ten do not.
        resp = _post(client, "form.docx", blank)
        assert resp.headers["X-EOI-Unmatched"] == "section-10"

    def test_filename_is_derived_from_the_upload_and_sanitised(
        self, client, blank
    ):
        resp = _post(client, "../../we*ird name.docx", blank)
        assert resp.status_code == 200
        assert (
            resp.headers["content-disposition"]
            == 'attachment; filename="weird name-filled.docx"'
        )

    def test_unrelated_docx_reports_every_rule(self, client):
        resp = _post(client, "letter.docx", _unrelated_docx())
        assert resp.status_code == 200
        unmatched = resp.headers["X-EOI-Unmatched"].split(",")
        assert unmatched == [f"section-{n}" for n in range(1, 12)]

    def test_rtf_upload_422(self, client):
        rtf = rb"{\rtf1\ansi Nothing here}"
        resp = _post(client, "form.rtf", rtf, "application/rtf")
        assert resp.status_code == 422
        assert ".docx" in resp.json()["detail"]

    def test_non_document_upload_422(self, client):
        resp = _post(client, "cert.pdf", b"%PDF-1.7\n%stub\n", "application/pdf")
        assert resp.status_code == 422

    def test_zip_that_is_not_a_docx_422(self, client):
        # Passes the leading-bytes check, fails when python-docx opens it.
        resp = _post(client, "form.docx", b"PK\x03\x04 not really a package")
        assert resp.status_code == 422

    def test_oversize_document_413(self, client):
        big = b"PK" + b"0" * (10 * 1024 * 1024 + 1)
        resp = _post(client, "form.docx", big)
        assert resp.status_code == 413
